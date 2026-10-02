import streamlit as st
import pandas as pd
import numpy as np
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Leads Converter & Updater", layout="wide")
st.title("📊 Leads Converter & Tracker Updater")

# ==========================================
# 1. AREA UPLOAD FILE
# ==========================================
col_up1, col_up2 = st.columns(2)
with col_up1:
    file_raw = st.file_uploader("1. Unggah Raw Data HQ Baru (.xlsx)", type=["xlsx"])
with col_up2:
    file_master_lama = st.file_uploader("2. Unggah Master Data Lama (Opsional)", type=["xlsx"])

st.markdown("---")

# ==========================================
# 2. FITUR MAPPING DINAMIS DI LAYAR DEPAN
# ==========================================
st.subheader("⚙️ 1. Mapping Struktur Data (Nama Sheet & Judul Kolom)")
st.info("Kamus pembaca data. Jika ada sheet baru atau nama kolom yang berubah dari pusat, tambahkan atau edit langsung di tabel bawah ini.")

# Default kamus yang sudah lengkap
default_mapping = pd.DataFrame({
    "Nama Sheet (Excel)": ["DPK All", "SME", "KPR", "payroll dpk", "Potensi Akuisisi Payroll", "Potensi Kredit SME_Comm", "Leads KPP Top Down", "casa vc", "kkb.", "CC", "ksm", "bod boc wealth", "KIPK"],
    "Kolom CIF": ["CIFNO", "CIF Number", "CIF", "cif_sumber", "CIF GAB", "CIF GAB", "No CIF", "CIF L2", "cifno", "CIF Perusahaan", "", "CIF", "CIF"],
    "Kolom Cabang": ["Cabang", "Branch", "Cabang", "branch_nm", "AREA", "AREA", "SME Banking Center", "Cabang L2", "area_name", "Nama Cabang", "hub_nm", "cabang_pengelola", "Cabang"],
    "Kolom Nama": ["SNAME", "Nama", "full_name", "nama pers", "RS", "RS", "Nama", "Nama L2", "nama", "Nama Principal", "nama_perusahaan", "NAMA LENGKAP NASABAH", "Nama"],
    "Kolom PIC": ["PIC", "PIC", "PIC", "PIC", "PIC", "PIC", "PIC ", "PIC", "PIC", "PIC", "Unnamed: 12", "PIC", "PIC"],
    "Kolom Potensi": ["End Balance 20.08'26", "Potensi Limit", "Potensi Limit", "potensi DPK", "TIDAK ADA", "TIDAK ADA", "Potensi Limit", "Potential Uplift Casa VC L2", "limit_indikatif", "Potensi Limit Indikatif", "Potensi", "POTENSI MINIMUM FUM", "Potensi"]
})
# Tabel interaktif untuk struktur
edit_mapping = st.data_editor(default_mapping, num_rows="dynamic", use_container_width=True, key="edit_mapping")


st.subheader("⚙️️ 2. Mapping Standarisasi Nama (Cabang & Produk)")
col_map1, col_map2 = st.columns(2)

with col_map1:
    st.write("**Ubah Kode Angka Menjadi Nama Cabang**")
    default_cabang = pd.DataFrame({
        "Nama Asli (Dari HQ)": [
            "13600", "13603", "13604", "13605", "13611", "13617", "13626", 
            "13628", "13612", "13619", "13618", "13621", "13622", "13625", 
            "13676", "13627", "13672", "13673", "13677", "13678", "13687", 
            "13688", "13689", "13695", "13696", "13699"
        ],
        "Ubah Menjadi": [
            "Semarang Pahlawan", "Semarang Candi Baru", "Semarang Sugiyopranoto", 
            "Ungaran", "Semarang Kawasan Industri Candi", "Semarang Universitas Diponegoro", 
            "Semarang Sisingamangaraja", "Salatiga", "Weleri", "Semarang Ngaliyan", 
            "Ambarawa", "Semarang Kedungmundu", "Semarang Puri Anjasmoro", 
            "Semarang Pamularsih", "Semarang Bandungan", "Kendal Soekarno Hatta", 
            "Tengaran", "Semarang Bringin", "Semarang Boja", "Kendal Sukorejo", 
            "Karangjati", "Kaliwungu", "Semarang UNNES", "Semarang BSB", 
            "Kendal Cepiring", "Semarang Manyaran"
        ]
    })
    edit_cabang = st.data_editor(default_cabang, num_rows="dynamic", use_container_width=True, key="edit_cabang")

with col_map2:
    st.write("**Ubah Nama Produk (Source) ke Nama Singkat**")
    default_produk = pd.DataFrame({
        "Nama Asli (Dari HQ)": ["Potensi Akuisisi Payroll", "Potensi Kredit SME_Comm", "payroll dpk", "DPK All", "Leads KPP Top Down"],
        "Ubah Menjadi": ["Leads RS", "Leads RS", "Payroll DPK", "DPK Reguler", "KPP Top Down"]
    })
    edit_produk = st.data_editor(default_produk, num_rows="dynamic", use_container_width=True, key="edit_produk")

st.markdown("---")


# ==========================================
# 3. FUNGSI UTAMA (KONVERSI & MERGE)
# ==========================================
def konversi_dan_sinkronisasi(raw_file, master_file, map_struktur, map_cabang, map_produk):
    xls = pd.ExcelFile(raw_file)
    
    # Menerjemahkan tabel struktur di layar menjadi dictionary untuk mesin
    mapping = {}
    for index, row in map_struktur.iterrows():
        sheet_name = str(row["Nama Sheet (Excel)"]).strip()
        if sheet_name and sheet_name != "nan" and sheet_name != "None":
            mapping[sheet_name] = {
                'cif': str(row["Kolom CIF"]).strip() if pd.notna(row["Kolom CIF"]) and str(row["Kolom CIF"]).strip() != "" else None,
                'cabang': str(row["Kolom Cabang"]).strip() if pd.notna(row["Kolom Cabang"]) and str(row["Kolom Cabang"]).strip() != "" else None,
                'nama': str(row["Kolom Nama"]).strip() if pd.notna(row["Kolom Nama"]) and str(row["Kolom Nama"]).strip() != "" else None,
                'pic': str(row["Kolom PIC"]).strip() if pd.notna(row["Kolom PIC"]) and str(row["Kolom PIC"]).strip() != "" else None,
                'potensi': str(row["Kolom Potensi"]).strip() if pd.notna(row["Kolom Potensi"]) and str(row["Kolom Potensi"]).strip() != "" else None
            }
            
    daftar_tabel = []
    for sheet, col in mapping.items():
        if sheet in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet)
            temp = pd.DataFrame()
            
            # Pengambilan data dengan validasi kolom
            temp['CIF'] = df[col['cif']].fillna("-").astype(str).str.replace(r'\.0$', '', regex=True) if col['cif'] and col['cif'] in df.columns else "-"
            temp['Cabang'] = df[col['cabang']].astype(str) if col['cabang'] and col['cabang'] in df.columns else "Unknown"
            temp['Nama Nasabah'] = df[col['nama']].astype(str) if col['nama'] and col['nama'] in df.columns else "Unknown"
            temp['PIC'] = df[col['pic']].astype(str) if col['pic'] and col['pic'] in df.columns else "-"
            temp['Potensi/Limit'] = pd.to_numeric(df[col['potensi']], errors='coerce').fillna(0) if col['potensi'] and col['potensi'] in df.columns else 0
            temp['Produk (Source)'] = sheet
            
            daftar_tabel.append(temp)
            
    if len(daftar_tabel) == 0:
        st.error("❌ Error: Tidak ada satupun nama sheet di file HQ yang cocok dengan Tabel Mapping Struktur.")
        st.info(f"Daftar sheet yang terdeteksi di file yang Anda unggah: {xls.sheet_names}")
        return None 
        
    df_baru = pd.concat(daftar_tabel, ignore_index=True)
    
    # Ganti nama cabang dan produk berdasarkan tabel edit di layar
    dict_cabang = dict(zip(map_cabang['Nama Asli (Dari HQ)'].astype(str), map_cabang['Ubah Menjadi'].astype(str)))
    dict_produk = dict(zip(map_produk['Nama Asli (Dari HQ)'].astype(str), map_produk['Ubah Menjadi'].astype(str)))
    
    df_baru['Cabang'] = df_baru['Cabang'].replace(dict_cabang)
    df_baru['Produk (Source)'] = df_baru['Produk (Source)'].replace(dict_produk)
    
    if master_file is not None:
        df_lama = pd.read_excel(master_file)
        catatan = df_lama[['CIF', 'Nama Nasabah', 'Status Kunjungan', 'Keterangan']].drop_duplicates(subset=['CIF', 'Nama Nasabah'])
        df_lama_base = df_lama.drop(columns=['Status Kunjungan', 'Keterangan'], errors='ignore')
        
        df_gabung = pd.concat([df_lama_base, df_baru], ignore_index=True)
        df_gabung = df_gabung.drop_duplicates(subset=['CIF', 'Nama Nasabah'], keep='last')
        
        df_hasil = pd.merge(df_gabung, catatan, on=['CIF', 'Nama Nasabah'], how='left')
    else:
        df_hasil = df_baru.copy()
        df_hasil['Status Kunjungan'] = ""
        df_hasil['Keterangan'] = ""
        
    df_hasil['Status Kunjungan'] = df_hasil['Status Kunjungan'].fillna("")
    df_hasil['Keterangan'] = df_hasil['Keterangan'].fillna("")
    
    kolom_final = ['CIF', 'Cabang', 'Nama Nasabah', 'PIC', 'Potensi/Limit', 'Produk (Source)', 'Status Kunjungan', 'Keterangan']
    df_hasil = df_hasil[kolom_final]
    
    return df_hasil


# ==========================================
# 4. TOMBOL EKSEKUSI & FORMATTING EXCEL
# ==========================================
if file_raw and st.button("🚀 Mulai Konversi & Sinkronisasi"):
    with st.spinner("Sedang memproses dan menyelaraskan data..."):
        # Parameter fungsi kini mencakup ketiga tabel edit dari layar
        hasil_df = konversi_dan_sinkronisasi(file_raw, file_master_lama, edit_mapping, edit_cabang, edit_produk)
        
        if hasil_df is not None:
            output = BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                hasil_df.to_excel(writer, index=False, sheet_name='Master_Data')
                
                worksheet = writer.sheets['Master_Data']
                header_font = Font(bold=True)
                header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid") 
                thin_border = Border(
                    left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin')
                )
                
                for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
                    for cell in row:
                        cell.border = thin_border 
                        if cell.row == 1: 
                            cell.font = header_font
                            cell.fill = header_fill
                
                for col in range(1, worksheet.max_column + 1):
                    worksheet.column_dimensions[get_column_letter(col)].width = 18
            
            st.success(f"✅ Berhasil mengolah {len(hasil_df):,} baris data gabungan!")
            st.download_button(
                label="📥 Unduh Master Data Terupdate",
                data=output.getvalue(),
                file_name="Master_Database_Kunjungan_Updated.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )